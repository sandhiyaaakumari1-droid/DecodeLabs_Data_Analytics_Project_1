"""
DecodeLabs | Data Analytics | Project 1: Data Cleaning & Preparation
Reproducible cleaning pipeline.

Usage:
    python clean_dataset.py <raw.xlsx> <cleaned_output.xlsx> <stats.json>

Phases (mirroring the training kit):
    Phase 1 - Strategic imputation (missing values)
    Phase 2 - Integrity audit (duplicates)
    Phase 3 - Speak one language (dates, text, numeric precision)
    Verification gate - zero duplicate IDs, zero badly formatted dates
"""
import sys
import json
from collections import Counter

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

RAW, OUT, STATS = sys.argv[1], sys.argv[2], sys.argv[3]

ID_COLS = ["OrderID", "CustomerID", "TrackingNumber"]
CATEGORICAL_TITLE = ["Product", "PaymentMethod", "OrderStatus", "ReferralSource"]
TEXT_COLS = ["OrderID", "CustomerID", "Product", "ShippingAddress", "PaymentMethod",
             "OrderStatus", "TrackingNumber", "CouponCode", "ReferralSource"]
MONEY_COLS = ["UnitPrice", "TotalPrice"]
INT_COLS = ["Quantity", "ItemsInCart"]
NO_COUPON = "No Coupon"

# ---------------------------------------------------------------- load raw
df = pd.read_excel(RAW)
n_raw = len(df)
stats = {"rows_raw": n_raw, "cols": list(df.columns)}

# raw cell-level inspection (pandas hides int-vs-float storage)
ws_raw = openpyxl.load_workbook(RAW).active
hdr = [c.value for c in ws_raw[1]]
int_stored = {}
for col in MONEY_COLS:
    j = hdr.index(col) + 1
    int_stored[col] = sum(isinstance(ws_raw.cell(r, j).value, int)
                          for r in range(2, ws_raw.max_row + 1))
stats["money_cells_stored_as_integer"] = int_stored

# ------------------------------------------------------- BEFORE-cleaning audit
stats["before"] = {
    "missing_by_column": {c: int(v) for c, v in df.isna().sum().items() if v > 0},
    "full_row_duplicates": int(df.duplicated().sum()),
    "duplicate_order_ids": int(df["OrderID"].duplicated().sum()),
    "duplicate_tracking_numbers": int(df["TrackingNumber"].duplicated().sum()),
    "near_duplicate_orders": int(df.duplicated(
        subset=[c for c in df.columns if c not in ("OrderID", "TrackingNumber")]).sum()),
}

# ------------------------------------------- PHASE 1: strategic imputation
missing_coupon_mask = df["CouponCode"].isna()
imputed_order_ids = df.loc[missing_coupon_mask, "OrderID"].tolist()
df["CouponCode"] = df["CouponCode"].fillna(NO_COUPON)
stats["imputed_coupon_rows"] = int(missing_coupon_mask.sum())
stats["imputed_coupon_pct"] = round(missing_coupon_mask.mean() * 100, 2)
stats["mode_coupon_would_have_been"] = df.loc[~missing_coupon_mask, "CouponCode"].mode()[0]

# --------------------------------------------- PHASE 2: integrity audit
before = len(df)
df = df.drop_duplicates(keep="first")                       # exact duplicate rows
df = df.drop_duplicates(subset="OrderID", keep="first")     # duplicate primary keys
stats["duplicate_rows_removed"] = before - len(df)

# --------------------------------------------- PHASE 3: one language
# 3a. Text: trim whitespace + collapse internal runs of spaces
ws_changes = 0
for c in TEXT_COLS:
    cleaned = df[c].astype(str).str.strip().str.replace(r"\s+", " ", regex=True)
    ws_changes += int((cleaned != df[c].astype(str)).sum())
    df[c] = cleaned
stats["whitespace_cells_fixed"] = ws_changes

# 3b. Proper Case on descriptive categoricals only (IDs / codes stay upper-case by design)
case_changes = 0
for c in CATEGORICAL_TITLE:
    cleaned = df[c].str.title()
    case_changes += int((cleaned != df[c]).sum())
    df[c] = cleaned
stats["case_cells_fixed"] = case_changes

# 3c. Dates -> ISO 8601 (YYYY-MM-DD), real date values (not text)
parsed = pd.to_datetime(df["Date"], errors="coerce")
stats["unparseable_dates"] = int(parsed.isna().sum())
stats["dates_with_time_component"] = int((parsed.dt.normalize() != parsed).sum())
df["Date"] = parsed.dt.normalize()

# 3d. Numerics
for c in INT_COLS:
    df[c] = pd.to_numeric(df[c], errors="raise").astype("int64")
float_artifacts = {}
for c in MONEY_COLS:
    num = pd.to_numeric(df[c], errors="raise")
    float_artifacts[c] = int((num != num.round(2)).sum())   # e.g. 769.3799999999999
    df[c] = num.round(2)
stats["float_artifact_cells_rounded"] = float_artifacts

# ------------------------------------------------- business-rule validation
calc = (df["Quantity"] * df["UnitPrice"]).round(2)
mismatch = (calc - df["TotalPrice"]).abs() > 0.005
stats["total_price_mismatches"] = int(mismatch.sum())

# flagged-but-retained observations (reviewed, valid -> no change)
q1, q3 = df["TotalPrice"].quantile([.25, .75])
iqr = q3 - q1
out = (df["TotalPrice"] < q1 - 1.5 * iqr) | (df["TotalPrice"] > q3 + 1.5 * iqr)
stats["totalprice_iqr_outliers_retained"] = int(out.sum())
stats["repeat_customer_ids_retained"] = int(df["CustomerID"].duplicated().sum())

df = df.sort_values("OrderID").reset_index(drop=True)
stats["rows_clean"] = len(df)

# ---------------------------------------------------- AFTER-cleaning proof (python)
stats["after"] = {
    "missing_cells_total": int(df.isna().sum().sum()),
    "duplicate_order_ids": int(df["OrderID"].duplicated().sum()),
    "duplicate_full_rows": int(df.duplicated().sum()),
    "non_date_values": int(df["Date"].isna().sum()),
    "dates_not_iso_yyyy_mm_dd": int((~df["Date"].dt.strftime("%Y-%m-%d")
                                     .str.match(r"^\d{4}-\d{2}-\d{2}$")).sum()),
    "date_min": df["Date"].min().strftime("%Y-%m-%d"),
    "date_max": df["Date"].max().strftime("%Y-%m-%d"),
}
stats["coupon_distribution"] = df["CouponCode"].value_counts().to_dict()

# ============================================================ write workbook
FONT = "Arial"
HDR_FILL = PatternFill("solid", start_color="1F3A5F")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True, color="1F3A5F")
NOTE = Font(name=FONT, size=9, italic=True, color="595959")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()

# ---- Sheet 1: Cleaned_Data
ws = wb.active
ws.title = "Cleaned_Data"
cols = list(df.columns)
ws.append(cols)
for rec in df.itertuples(index=False):
    row = []
    for c, v in zip(cols, rec):
        if c == "Date":
            v = v.to_pydatetime()
        elif c in INT_COLS:
            v = int(v)
        elif c in MONEY_COLS:
            v = float(v)
        row.append(v)
    ws.append(row)

n = len(df)
last = n + 1
for j, c in enumerate(cols, start=1):
    h = ws.cell(1, j)
    h.font, h.fill, h.border = HDR_FONT, HDR_FILL, BOX
    h.alignment = Alignment(horizontal="center", vertical="center")
    for r in range(2, last + 1):
        cell = ws.cell(r, j)
        cell.font = BODY
        if c == "Date":
            cell.number_format = "yyyy-mm-dd"
            cell.alignment = Alignment(horizontal="center")
        elif c in MONEY_COLS:
            cell.number_format = "#,##0.00"
        elif c in INT_COLS:
            cell.number_format = "0"
    width = max(len(c), *(len(str(ws.cell(r, j).value)) for r in range(2, min(last, 60) + 1))) + 3
    ws.column_dimensions[get_column_letter(j)].width = min(max(width, 11), 22)
ws.row_dimensions[1].height = 22
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"

# column letters for formulas
L = {c: get_column_letter(i + 1) for i, c in enumerate(cols)}
rng = lambda c: f"Cleaned_Data!${L[c]}$2:${L[c]}${last}"
full = f"Cleaned_Data!$A$2:${get_column_letter(len(cols))}${last}"

# ---- Sheet 2: Verification (live formulas)
v = wb.create_sheet("Verification")
v["A1"] = "Verification Gate: Threshold for Project 2"
v["A1"].font = TITLE
v["A2"] = ('"Before you finish, you must prove there are zero duplicate IDs and '
           'zero incorrectly formatted dates."')
v["A2"].font = NOTE

heads = ["#", "Check", "Rule", "Result (live formula)", "Target", "Status"]
for j, h in enumerate(heads, start=1):
    c = v.cell(4, j, h)
    c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BOX
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

checks = [
    ("Duplicate OrderIDs", "Every OrderID appears exactly once",
     f"=SUMPRODUCT(--(COUNTIF({rng('OrderID')},{rng('OrderID')})>1))", 0),
    ("Duplicate TrackingNumbers", "Every TrackingNumber appears exactly once",
     f"=SUMPRODUCT(--(COUNTIF({rng('TrackingNumber')},{rng('TrackingNumber')})>1))", 0),
    ("Non-date values in Date column", "Every Date cell is a true date value (not text)",
     f"=SUMPRODUCT(--NOT(ISNUMBER({rng('Date')})))", 0),
    ("Dates with a time component", "Every Date is a whole day (YYYY-MM-DD)",
     f"=SUMPRODUCT(--({rng('Date')}<>INT({rng('Date')})))", 0),
    ("Dates outside expected window", "Dates fall between 2023-01-01 and 2025-06-30",
     f"=SUMPRODUCT(--(({rng('Date')}<DATE(2023,1,1))+({rng('Date')}>DATE(2025,6,30))>0))", 0),
    ("Blank cells in dataset", "No missing values remain in any column",
     f"=COUNTBLANK({full})", 0),
    ("TotalPrice arithmetic errors", "TotalPrice = Quantity x UnitPrice (to the cent)",
     f"=SUMPRODUCT(--(ABS({rng('Quantity')}*{rng('UnitPrice')}-{rng('TotalPrice')})>0.005))", 0),
    ("Cells with stray whitespace", "No leading/trailing/double spaces in any cell",
     f"=SUMPRODUCT(--(LEN({full})<>LEN(TRIM({full}))))", 0),
    ("Non-positive Quantity / UnitPrice", "All quantities and prices are > 0",
     f"=SUMPRODUCT(--({rng('Quantity')}<=0))+SUMPRODUCT(--({rng('UnitPrice')}<=0))", 0),
]
r = 5
for i, (name, rule, formula, target) in enumerate(checks, start=1):
    v.cell(r, 1, i)
    v.cell(r, 2, name)
    v.cell(r, 3, rule)
    v.cell(r, 4, formula)
    v.cell(r, 5, target)
    v.cell(r, 6, f'=IF(IFERROR(D{r}=E{r},FALSE),"PASS","FAIL")')
    for j in range(1, 7):
        c = v.cell(r, j)
        c.font, c.border = BODY, BOX
        c.alignment = Alignment(horizontal="center" if j in (1, 4, 5, 6) else "left",
                                vertical="center", wrap_text=True)
    v.cell(r, 6).font = BOLD
    r += 1
last_check = r - 1

r += 1
v.cell(r, 2, "Record count (rows in Cleaned_Data)").font = BOLD
v.cell(r, 4, f"=COUNTA({rng('OrderID')})").font = BOLD
v.cell(r, 4).alignment = Alignment(horizontal="center")
r += 1
v.cell(r, 2, "Unique OrderIDs").font = BOLD
v.cell(r, 4, f"=SUMPRODUCT(1/COUNTIF({rng('OrderID')},{rng('OrderID')}))").font = BOLD
v.cell(r, 4).alignment = Alignment(horizontal="center")
r += 2
v.cell(r, 2, "OVERALL GATE").font = Font(name=FONT, size=12, bold=True)
v.cell(r, 4, f'=IF(COUNTIF(F5:F{last_check},"FAIL")=0,"PASS: eligible for Project 2","FAIL: review checks above")')
v.cell(r, 4).font = Font(name=FONT, size=12, bold=True, color="006100")
v.cell(r, 4).fill = PatternFill("solid", start_color="C6EFCE")
v.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
v.cell(r, 4).alignment = Alignment(horizontal="center")
gate_row = r

r += 2
v.cell(r, 2, "Baseline recorded from the ORIGINAL file (static values, captured by clean_dataset.py)").font = BOLD
r += 1
base = [
    ("Rows in raw file", n_raw),
    ("Missing values in raw file (all in CouponCode)", stats["before"]["missing_by_column"].get("CouponCode", 0)),
    ("Duplicate OrderIDs in raw file", stats["before"]["duplicate_order_ids"]),
    ("Duplicate rows removed", stats["duplicate_rows_removed"]),
    ("Rows in cleaned file", stats["rows_clean"]),
]
for label, val in base:
    v.cell(r, 2, label).font = BODY
    c = v.cell(r, 4, val)
    c.font = Font(name=FONT, size=10, color="0000FF")  # blue = hardcoded input
    c.alignment = Alignment(horizontal="center")
    r += 1
r += 1
v.cell(r, 2, "Note: blue = static baseline from the original file; black = live formula against Cleaned_Data.").font = NOTE

for col, w in zip("ABCDEF", [5, 38, 52, 22, 10, 12]):
    v.column_dimensions[col].width = w
v.sheet_view.showGridLines = False

# ---- Sheet 3: Imputed_Rows (traceability for CR001)
im = wb.create_sheet("Imputed_Rows")
im["A1"] = "Rows where blank CouponCode was set to 'No Coupon' (change CR001)"
im["A1"].font = Font(name=FONT, size=12, bold=True, color="1F3A5F")
for j, h in enumerate(["OrderID", "Original CouponCode", "Cleaned CouponCode"], start=1):
    c = im.cell(3, j, h)
    c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BOX
    c.alignment = Alignment(horizontal="center")
for i, oid in enumerate(sorted(imputed_order_ids), start=4):
    im.cell(i, 1, oid)
    im.cell(i, 2, "(blank)")
    im.cell(i, 3, NO_COUPON)
    for j in range(1, 4):
        im.cell(i, j).font = BODY
for col, w in zip("ABC", [16, 22, 22]):
    im.column_dimensions[col].width = w
im.freeze_panes = "A4"

wb.save(OUT)
stats["verification_gate_row"] = gate_row
with open(STATS, "w") as f:
    json.dump(stats, f, indent=2, default=str)
print(json.dumps(stats, indent=2, default=str))
